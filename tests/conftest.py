from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import Workbook

FIXTURES = Path(__file__).parent / "fixtures"

# Pääikkunan automaattinen päivitystarkistus ei saa tehdä verkkokutsuja testeissä
os.environ.setdefault("SUUNNITELMOITTAJA_NO_UPDATE_CHECK", "1")


@pytest.fixture
def excel_file(tmp_path: Path) -> Path:
    """Pieni Excel, jossa on kaksi kelvollista välilehteä ja yksi rikkinäinen."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ohjelmistokehittäjä"
    ws.append(
        ["Osaamistavoite", "Laajuus", "Suoritustapa / osaaminen hankitaan", "Suoritusajankohta"]
    )
    ws.append(["Ohjelmointi", 45, "Joustava - itsLearning", None])
    ws.append(["Tietoturva", 30.0, "Lähiopetus", "syksy 2026"])
    ws.append([None, None, None, None])  # tyhjä rivi ohitetaan

    ws2 = wb.create_sheet("Kyber")
    ws2.append(
        ["Osaamistavoite", "Laajuus", "Suoritustapa / osaaminen hankitaan", "Suoritusajankohta"]
    )
    ws2.append(["Kyberturvallisuuden ylläpitäminen", 15.5, "Verkko", None])

    ws3 = wb.create_sheet("Rikki")
    ws3.append(["Osaamistavoite", "Laajuus"])
    ws3.append(["Jotain", 5])

    path = tmp_path / "testi.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def lomake_url() -> str:
    return (FIXTURES / "lomake.html").resolve().as_uri()


@pytest.fixture
def wilma_url() -> str:
    """Wilman lomakkeen rakennekopio: samat id-kaavat, neljä tallennettua riviä, tfootin nappi."""
    return (FIXTURES / "wilma-lomake.html").resolve().as_uri()


@pytest.fixture
def tallennettu_url() -> str:
    """Wilman tallennettua lomaketta jäljittelevä sivu: rivit ilman poistonappia."""
    return (FIXTURES / "lomake_tallennettu.html").resolve().as_uri()


@pytest.fixture(scope="session")
def browser_type_launch_args(browser_type_launch_args: dict[str, object]) -> dict[str, object]:
    """Salli Chromium-binäärin osoittaminen ympäristömuuttujalla (CI, valmis selain)."""
    override = os.environ.get("SUUNNITELMOITTAJA_CHROMIUM")
    if override:
        return {**browser_type_launch_args, "executable_path": override}
    return browser_type_launch_args
