"""Excel-tiedoston luku ilman pandasia (openpyxl riittää)."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

log = logging.getLogger("suunnitelmoittaja.excel")


class ExcelError(Exception):
    """Excel-tiedostoa ei voi lukea tai sen rakenne ei kelpaa."""


@dataclass(frozen=True, slots=True)
class PlanRow:
    """Yksi lomakkeelle vietävä rivi: kenttänimi → arvo merkkijonona (tyhjä = ei arvoa)."""

    values: dict[str, str]

    def is_empty(self) -> bool:
        return not any(v.strip() for v in self.values.values())


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    rows: list[PlanRow]


def list_sheets(path: Path) -> list[str]:
    wb = _open(path)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_sheet(path: Path, sheet_name: str, columns: dict[str, str]) -> Sheet:
    """Lue yksi välilehti. ``columns`` kuvaa kenttänimen → Excelin otsikkosolun tekstiin."""
    wb = _open(path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ExcelError(f"Välilehteä '{sheet_name}' ei ole tiedostossa {path}")
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ExcelError(f"Välilehti '{sheet_name}' on tyhjä")

        header_index = {_norm(h): i for i, h in enumerate(header) if h is not None}
        missing = [title for title in columns.values() if _norm(title) not in header_index]
        if missing:
            raise ExcelError(
                f"Välilehdeltä '{sheet_name}' puuttuvat sarakkeet: {missing}. "
                f"Löydetyt otsikot: {[h for h in header if h is not None]}"
            )

        rows: list[PlanRow] = []
        for raw in iterator:
            values = {
                field: format_cell(raw[header_index[_norm(title)]])
                for field, title in columns.items()
            }
            row = PlanRow(values)
            if not row.is_empty():
                rows.append(row)
        log.info("Välilehti '%s': %d riviä", sheet_name, len(rows))
        return Sheet(sheet_name, rows)
    finally:
        wb.close()


def resolve_sheet_selection(selection: str, available: list[str]) -> list[str]:
    """Muuta käyttäjän syöte ("1,3,5" tai välilehtien nimiä) välilehtien nimiksi.

    Numerot ovat 1-pohjaisia. Kelpaamattomat ohitetaan; jos mikään ei kelpaa,
    nostetaan ValueError.
    """
    chosen: list[str] = []
    for token in (t.strip() for t in selection.split(",")):
        if not token:
            continue
        if token.isdigit():
            idx = int(token) - 1
            if 0 <= idx < len(available):
                chosen.append(available[idx])
                continue
        elif token in available:
            chosen.append(token)
            continue
        log.warning("Ohitetaan kelpaamaton välilehtivalinta: %r", token)
    if not chosen:
        raise ValueError(f"Yhtään kelvollista välilehteä ei valittu syötteestä {selection!r}")
    return chosen


def format_cell(value: Any) -> str:
    """Muuta solun arvo lomakkeelle sopivaksi tekstiksi (25.0 → "25", None → "")."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "kyllä" if value else "ei"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y") if value.time() == datetime.min.time() else str(value)
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _open(path: Path):
    if not path.exists():
        raise ExcelError(f"Excel-tiedostoa ei löydy: {path}")
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl nostaa sekalaisia poikkeuksia
        raise ExcelError(f"Excel-tiedoston {path} avaaminen epäonnistui: {exc}") from exc


def _norm(title: Any) -> str:
    return " ".join(str(title).split()).casefold()
